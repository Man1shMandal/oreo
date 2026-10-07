"""Files uploaded from the browser: documents become text, images are kept for the model to see.

The ABB gateway only takes pictures as a PDF in an `input_file` part, and only Claude models read
them, so images are wrapped in a one-page PDF and scanned PDFs are sent as they are."""

import base64
import hashlib
import io
from pathlib import Path

MAX_UPLOAD = 15_000_000     # bytes per file
MAX_TEXT = 200_000          # characters of extracted text per file
MAX_SIDE = 1568             # bigger images are scaled down (fewer tokens, same answer)
IMAGES = {"image/png": "png", "image/jpeg": "jpg", "image/webp": "webp", "image/gif": "gif"}


def read_document(name, data):
    ext = Path(name).suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader
        pages = [p.extract_text() or "" for p in PdfReader(io.BytesIO(data)).pages]
        if not "".join(pages).strip():
            return ""   # scanned, no text layer
        return "\n\n".join(f"--- page {i} ---\n{t}" for i, t in enumerate(pages, 1))
    if ext == ".docx":
        import docx
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            parts += [" | ".join(c.text for c in row.cells) for row in t.rows]
        return "\n".join(parts)
    try:
        return data.decode()
    except UnicodeDecodeError:
        raise ValueError("can't read this kind of file (try PDF, Word, text or an image)")


def save_image(data, folder):
    """Shrink if needed, store under its hash, return the file name."""
    from PIL import Image
    img = Image.open(io.BytesIO(data))
    fmt = "PNG" if img.format in ("PNG", "GIF") else "JPEG"
    if max(img.size) > MAX_SIDE or img.format not in ("PNG", "JPEG"):
        img.thumbnail((MAX_SIDE, MAX_SIDE))
        out = io.BytesIO()
        (img if fmt == "PNG" else img.convert("RGB")).save(out, fmt, quality=88)
        data = out.getvalue()
    name = hashlib.sha256(data).hexdigest()[:24] + (".png" if fmt == "PNG" else ".jpg")
    folder.mkdir(parents=True, exist_ok=True)
    (folder / name).write_bytes(data)
    return name


def save_pdf(data, folder):
    name = hashlib.sha256(data).hexdigest()[:24] + ".pdf"
    folder.mkdir(parents=True, exist_ok=True)
    (folder / name).write_bytes(data)
    return name


def process(uploads, folder):
    """uploads: [{name, type, data(base64)}] -> (text blocks, image names, errors)."""
    blocks, images, errors = [], [], []
    for u in uploads or []:
        name = Path(str(u.get("name") or "file")).name.replace('"', "")
        try:
            data = base64.b64decode(u.get("data") or "")
            if len(data) > MAX_UPLOAD:
                raise ValueError("too big (over 15 MB)")
            if u.get("type") in IMAGES:
                images.append(save_image(data, folder))
                continue
            text = read_document(name, data).strip()
            if not text and name.lower().endswith(".pdf"):   # scanned: let the model look at the pages
                images.append(save_pdf(data, folder))
                blocks.append(f'<file path="{name}">\n[scanned PDF, sent as pages]\n</file>')
                continue
            if not text:
                raise ValueError("no text found")
            if len(text) > MAX_TEXT:
                text = text[:MAX_TEXT] + "\n[…cut, file too long]"
            blocks.append(f'<file path="{name}">\n{text}\n</file>')
        except Exception as e:
            errors.append(f"{name}: {e}")
    return blocks, images, errors


def image_part(folder, name):
    data = (folder / name).read_bytes()
    if not name.endswith(".pdf"):
        from PIL import Image
        out = io.BytesIO()
        Image.open(io.BytesIO(data)).convert("RGB").save(out, "PDF")
        data = out.getvalue()
    return {"type": "input_file", "file_data": base64.b64encode(data).decode(), "filename": "attachment.pdf"}
